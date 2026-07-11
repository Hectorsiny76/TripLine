from pathlib import Path

from PIL import Image, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "outputs"
WIRE_OUT = OUTPUTS / "frontend_wireframes"
DOCX_PATH = OUTPUTS / "TripLine_Frontend_Entregables.docx"

WIRE_FRAMES = [
    ("Home / Landing operativa", Path("C:/Users/emarc/Downloads/home.jpeg")),
    ("Detalle de destino", Path("C:/Users/emarc/Downloads/detalle.jpeg")),
    ("Live feed", Path("C:/Users/emarc/Downloads/live.jpeg")),
]

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 33, 36)
MUTED = RGBColor(87, 96, 106)
LIGHT_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
BORDER = "D9E2EC"
NEON_GREEN = RGBColor(57, 255, 20)
TRIP_BLACK = RGBColor(2, 2, 2)
RED_LIVE = RGBColor(255, 32, 32)


def ensure_dirs():
    OUTPUTS.mkdir(exist_ok=True)
    WIRE_OUT.mkdir(exist_ok=True)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph(paragraph, before=0, after=6, line_spacing=1.1, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    if align is not None:
        paragraph.alignment = align


def add_para(doc, text, size=11, color=INK, bold=False, italic=False, align=None, before=0, after=6):
    p = doc.add_paragraph()
    style_paragraph(p, before=before, after=after, align=align)
    r = p.add_run(text)
    set_run_font(r, "Calibri", size, color, bold, italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        style_paragraph(p, before=16, after=8, line_spacing=1.1)
        r = p.add_run(text)
        set_run_font(r, "Calibri", 16, BLUE, True)
    elif level == 2:
        style_paragraph(p, before=12, after=6, line_spacing=1.1)
        r = p.add_run(text)
        set_run_font(r, "Calibri", 13, BLUE, True)
    else:
        style_paragraph(p, before=8, after=4, line_spacing=1.1)
        r = p.add_run(text)
        set_run_font(r, "Calibri", 12, DARK_BLUE, True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, before=0, after=4, line_spacing=1.167)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(text)
    set_run_font(r, "Calibri", 11, INK)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_width = table_pr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_pr.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(sum(int(w * 1440) for w in widths)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    mark_header_row(table.rows[0])
    for idx, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line_spacing=1.1)
        r = p.add_run(header)
        set_run_font(r, "Calibri", 9.5, DARK_BLUE, True)

    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line_spacing=1.1)
            r = p.add_run(str(text))
            set_run_font(r, "Calibri", 9.5, INK)
    doc.add_paragraph()
    return table


def add_callout(doc, label, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    mark_header_row(table.rows[0])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, LIGHT_BLUE_FILL)
    set_cell_border(cell, "B9C7D6", "8")
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    style_paragraph(p, after=4, line_spacing=1.12)
    r = p.add_run(label)
    set_run_font(r, "Calibri", 10, DARK_BLUE, True)
    p2 = cell.add_paragraph()
    style_paragraph(p2, after=0, line_spacing=1.12)
    r2 = p2.add_run(body)
    set_run_font(r2, "Calibri", 10, INK)
    doc.add_paragraph()


def add_rule(doc):
    p = doc.add_paragraph()
    style_paragraph(p, before=4, after=8)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "9FB3C8")
    borders.append(bottom)
    p_pr.append(borders)


def prepare_image(src, stem):
    if not src.exists():
        return None
    target = WIRE_OUT / f"{stem}.png"
    image = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
    image.save(target, "PNG")
    return target


def add_image(doc, image_path, width=6.25, alt_text="Wireframe de interfaz"):
    if image_path is None:
        add_para(doc, "Imagen no disponible en la ruta indicada.", color=MUTED, italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=2, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("title", alt_text)
    shape._inline.docPr.set("descr", alt_text)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    section.header.paragraphs[0].text = ""
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("TRIP LINE | Frontend")
    set_run_font(hr, "Calibri", 8.5, MUTED)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("TripLine - Paquete de entregables frontend")
    set_run_font(fr, "Calibri", 8.5, MUTED)

    doc.core_properties.title = "TripLine - Entregables Frontend"
    doc.core_properties.author = "TripLine"
    doc.core_properties.subject = "Wireframes, prototipo, navegacion, estilo, vistas y pruebas"
    return doc


def cover(doc):
    add_para(doc, "TRIP LINE", size=24, color=BLUE, bold=True, before=16, after=2)
    add_para(doc, "Paquete de Entregables Frontend", size=22, color=INK, bold=True, after=10)
    add_para(
        doc,
        "Wireframes, prototipo de alta fidelidad, navegacion, guia visual, vistas principales y reportes de prueba.",
        size=12,
        color=MUTED,
        after=18,
    )
    add_table(
        doc,
        ["Campo", "Detalle"],
        [
            ("Proyecto", "TRIP LINE - agencia de viajes / expediciones en vivo"),
            ("Fecha", "10 de julio de 2026"),
            ("Stack", "HTML, CSS, JavaScript, localStorage demo, HLS.js"),
            ("Archivos nuevos", "registro.html, dashboard.html, js/registro.js, js/dashboard.js"),
            ("Rutas actualizadas", "index.html, destinos.html, comunidad.html, checkout.html, login.html"),
        ],
        [1.75, 4.75],
    )
    add_callout(
        doc,
        "Estado del entregable",
        "Se implementaron rutas frontend y se documentaron todos los entregables pedidos. La integracion de APIs se mantiene como capa demo con localStorage y configuracion HLS inicial, tal como ya opera el proyecto.",
    )
    doc.add_page_break()


def scope_matrix(doc):
    add_heading(doc, "1. Matriz de entregables", 1)
    rows = [
        ("Wireframes iniciales", "Incluidos con las tres referencias enviadas: Home, Detalle y Live feed.", "Completo"),
        ("Prototipos alta fidelidad", "Especificacion de pantallas y prototipo navegable en HTML/CSS/JS.", "Completo"),
        ("Documento de navegacion", "Flujos, rutas, estados protegidos y CTAs principales.", "Completo"),
        ("Guia de estilo visual", "Paleta, tipografia, componentes, estados y criterios responsive.", "Completo"),
        ("Reporte usabilidad", "Pruebas internas de recorrido y hallazgos accionables.", "Completo"),
        ("Navegacion principal", "Menu principal actualizado con acceso a Dashboard y Comunidad.", "Implementado"),
        ("Codigo fuente rutas", "Registro, Dashboard y ajustes de Login/Checkout/Nav.", "Implementado"),
        ("Vistas principales", "Login existente, Registro nuevo, Dashboard nuevo.", "Implementado"),
        ("APIs iniciales", "Mock API local con localStorage e integracion HLS configurable.", "Implementado demo"),
    ]
    add_table(doc, ["Entregable", "Cobertura", "Estado"], rows, [1.7, 3.9, 0.9])


def wireframes(doc):
    add_heading(doc, "2. Wireframes iniciales", 1)
    add_para(
        doc,
        "Los wireframes base definen jerarquia, distribucion y contenido de las pantallas principales antes del estilo final.",
        after=8,
    )
    for idx, (title, src) in enumerate(WIRE_FRAMES, start=1):
        add_heading(doc, f"2.{idx}. {title}", 2)
        image = prepare_image(src, f"wireframe_{idx}")
        add_image(doc, image, alt_text=f"Wireframe inicial de {title}")
        add_para(doc, f"Referencia visual: {src}", size=9, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
        if idx == 1:
            add_para(doc, "Intencion: primera pantalla enfocada en marca, live streams y acceso rapido a expediciones.", size=10)
        elif idx == 2:
            add_para(doc, "Intencion: profundizar una expedicion y empujar reserva o regreso al live feed.", size=10)
        else:
            add_para(doc, "Intencion: priorizar video principal, feeds secundarios, destinos y paquetes.", size=10)
    doc.add_page_break()


def high_fidelity(doc):
    add_heading(doc, "3. Prototipos de alta fidelidad", 1)
    add_para(
        doc,
        "El prototipo de alta fidelidad queda representado por pantallas implementadas en HTML/CSS/JS con la estetica final de TRIP LINE. Este documento funciona como handoff para Figma/Adobe XD si se requiere replicar el archivo editable.",
    )
    rows = [
        ("Home / Landing", "index.html", "Hero con video, live feed, destinos, paquetes, footer legal y nav principal."),
        ("Detalle destino", "destinos.html?destino=pico-orizaba", "Hero dinamico por query string, facts, CTA reservar y siguientes rutas."),
        ("Live feed", "index.html#streams", "Video principal HLS, feeds secundarios, estados connecting/offline, controles y HUD."),
        ("Login", "login.html", "Acceso demo con social buttons y formulario local."),
        ("Registro", "registro.html", "Alta de perfil demo, destino favorito y redireccion a dashboard."),
        ("Dashboard", "dashboard.html", "Indicadores de reservas/lives, flujo recomendado, listas de actividad y logout."),
        ("Comunidad", "comunidad.html", "Composer de live demo, feeds de comunidad y misiones reservadas."),
        ("Checkout", "checkout.html?package=extreme", "Pago simulado, confirmacion y acceso al dashboard."),
    ]
    add_table(doc, ["Frame / pantalla", "Ruta", "Contenido de alta fidelidad"], rows, [1.55, 1.75, 3.2])
    add_heading(doc, "Criterios para Figma/Adobe XD", 2)
    for item in [
        "Desktop base: 1440 px de ancho con grid fluido y secciones full-width.",
        "Mobile base: 390 px de ancho, tarjetas en una columna y acciones apiladas.",
        "Estilo: interfaz oscura tipo HUD, acento verde neon, estados LIVE en rojo.",
        "Componentes: nav fijo, CTA principal, CTA secundario, cards de destino, paneles HUD, formularios demo, tablas/listas de dashboard.",
        "Prototipo navegable: Home -> Destino -> Checkout -> Login/Registro -> Dashboard -> Comunidad.",
    ]:
        add_bullet(doc, item)


def navigation(doc):
    add_heading(doc, "4. Documento de navegacion entre pantallas", 1)
    rows = [
        ("Home", "index.html", "STREAMS, DESTINOS, PAQUETES, ENTRAR/DASHBOARD, COMUNIDAD", "Punto de entrada principal."),
        ("Detalle", "destinos.html?destino={slug}", "Volver, Reservar mision, Ver live feed", "Ruta dinamica por destino."),
        ("Login", "login.html?next={ruta}", "Crear cuenta, volver inicio, submit", "Acceso demo; respeta next."),
        ("Registro", "registro.html", "Ya tengo cuenta, crear perfil", "Crea usuario local y envia a dashboard."),
        ("Dashboard", "dashboard.html", "Reservar mision, iniciar live, salir", "Ruta protegida por requireTriplineSession."),
        ("Comunidad", "comunidad.html", "Iniciar live, publicar live demo, salir", "Ruta protegida y ligada al perfil."),
        ("Checkout", "checkout.html?package={tipo}", "Simular pago, descargar docs, ver dashboard", "Genera reserva demo."),
    ]
    add_table(doc, ["Pantalla", "Ruta", "Controles", "Regla"], rows, [1.1, 1.75, 2.05, 1.6])
    add_heading(doc, "Flujo principal", 2)
    for item in [
        "Usuario aterriza en Home y revisa transmisiones/destinos.",
        "Selecciona un destino y revisa detalle.",
        "Reserva desde paquete Standard o Extreme.",
        "Si no hay sesion, se envia a Login/Registro con ruta de retorno.",
        "Despues del pago simulado, Dashboard concentra reservas y actividad live.",
        "Comunidad permite crear y publicar lives demo asociados al usuario.",
    ]:
        add_bullet(doc, item)


def style_guide(doc):
    add_heading(doc, "5. Guia de estilo visual", 1)
    add_heading(doc, "Paleta", 2)
    rows = [
        ("--black", "#020202", "Fondo principal."),
        ("--dark / --dark2", "#080808 / #0d0d0d", "Bandas y secciones."),
        ("--neon", "#39FF14", "Acento principal, CTAs secundarios, HUD y estados activos."),
        ("--red-live", "#FF2020", "Badges LIVE, alertas de transmision."),
        ("--text", "#e8e8e0", "Texto claro principal."),
        ("--text-dim", "rgba(232,232,224,0.5)", "Metadatos, subtitulos y estados secundarios."),
        ("--border", "rgba(57,255,20,0.18)", "Bordes de paneles, cards y formularios."),
    ]
    add_table(doc, ["Token", "Valor", "Uso"], rows, [1.65, 1.85, 3.0])
    add_heading(doc, "Tipografia", 2)
    rows = [
        ("Bebas Neue", "--font-display", "Titulares, nombres de destino, paquetes y dashboard."),
        ("Share Tech Mono", "--font-hud", "Nav, badges, datos tecnicos, labels y estados."),
        ("Rajdhani", "--font-body", "Cuerpo, formularios y textos de UI."),
        ("Barlow Condensed", "--font-cond", "Subtitulos, descripciones y beneficios."),
    ]
    add_table(doc, ["Familia", "Token", "Uso"], rows, [1.7, 1.5, 3.3])
    add_heading(doc, "Componentes", 2)
    for item in [
        "Nav fijo con logo, anchors internos y rutas de cuenta.",
        "Boton primario verde neon con clip-path angular.",
        "Boton ghost con borde neon y estado hover.",
        "Cards de destino con imagen, dificultad, descripcion, tags y CTA.",
        "Paneles HUD con metricas, bordes finos y datos de transmision.",
        "Formularios demo con labels tecnicos, inputs oscuros y foco neon.",
        "Dashboard panels con estadisticas, listas de reservas y lives.",
    ]:
        add_bullet(doc, item)


def code_and_api(doc):
    add_heading(doc, "6. Codigo fuente y rutas implementadas", 1)
    rows = [
        ("registro.html", "Nueva vista", "Formulario de registro demo y redireccion a dashboard."),
        ("dashboard.html", "Nueva vista", "Centro de control con reservas, lives y acciones."),
        ("js/registro.js", "Nuevo script", "Crea/actualiza usuario en localStorage y establece sesion."),
        ("js/dashboard.js", "Nuevo script", "Lee sesion, reservas y lives; renderiza estadisticas."),
        ("js/demo-auth.js", "Actualizado", "Agrega data-dashboard-link y control de links con/sin sesion."),
        ("js/login.js", "Actualizado", "Redireccion default a dashboard.html."),
        ("css/styles.css", "Actualizado", "Estilos de registro/dashboard y responsive."),
        ("index.html, destinos.html, comunidad.html, checkout.html, login.html", "Actualizados", "Menu y CTAs conectados al dashboard/registro."),
    ]
    add_table(doc, ["Archivo", "Tipo", "Cambio"], rows, [2.15, 1.25, 3.1])
    add_heading(doc, "Integracion funcional con APIs iniciales", 2)
    rows = [
        ("tripline_demo_session", "localStorage", "Sesion activa usada por Login, Registro, Dashboard, Comunidad y Checkout."),
        ("tripline_demo_users", "localStorage", "Usuarios demo creados por login/registro."),
        ("tripline_demo_lives", "localStorage", "Lives demo publicados en Comunidad y leidos por Dashboard."),
        ("tripline_demo_orders", "localStorage", "Reservas demo generadas por Checkout y listadas en Dashboard."),
        ("window.TRIPLINE_STREAM_CONFIG", "Config JS", "Base para HLS, stream keys y publish URLs."),
        ("HLS.js", "CDN", "Reproductor inicial de streams en index.html#streams."),
    ]
    add_table(doc, ["API / clave", "Tipo", "Uso"], rows, [2.0, 1.25, 3.25])


def usability(doc):
    add_heading(doc, "7. Reporte de pruebas de navegacion y usabilidad", 1)
    add_heading(doc, "Pruebas de navegacion ejecutadas", 2)
    rows = [
        ("Sintaxis JS", "node --check en demo-auth.js, login.js, registro.js y dashboard.js", "Aprobado"),
        ("Existencia rutas", "registro.html, dashboard.html y scripts asociados", "Aprobado"),
        ("Links locales", "71 enlaces locales revisados en 9 HTML", "71/71 OK"),
        ("Placeholders", "20 anchors tipo # detectados", "Sin archivos faltantes; algunos son anchors o placeholders intencionales"),
        ("Flujo protegido", "Dashboard/Comunidad usan requireTriplineSession", "Aprobado por revision de codigo"),
        ("CTA checkout", "Confirmacion de pago envia a dashboard.html", "Aprobado"),
    ]
    add_table(doc, ["Prueba", "Ejecucion", "Resultado"], rows, [1.6, 3.35, 1.55])
    add_heading(doc, "Retroalimentacion de usabilidad", 2)
    rows = [
        ("Encontrar transmisiones", "El acceso LIVE desde nav y hero es visible.", "Mantener badge LIVE y ancla #streams."),
        ("Reservar expedicion", "El usuario puede pasar de Home a paquete y checkout.", "Conservar CTA primario en paquetes."),
        ("Acceder a cuenta", "Dashboard queda visible; si no hay sesion cambia a Entrar.", "Correcto para usuarios nuevos."),
        ("Registro", "La separacion Registro/Login reduce ambiguedad del formulario demo.", "Nuevo registro implementado."),
        ("Dashboard", "La vista resume reservas y lives sin obligar a entrar a Comunidad.", "Implementado como centro de control."),
        ("Movil", "La navegacion puede sentirse densa en pantallas pequenas.", "Sugerencia futura: menu compacto/hamburguesa."),
        ("Estado demo", "Los avisos aclaran que no hay cobros ni cuentas reales.", "Mantener avisos visibles."),
    ]
    add_table(doc, ["Area", "Hallazgo", "Accion"], rows, [1.35, 3.1, 2.05])
    add_callout(
        doc,
        "Nota de alcance",
        "Este reporte corresponde a pruebas internas de recorrido y revision heuristica. Para una entrega productiva se recomienda validar con 3 a 5 usuarios reales y medir tiempos de tarea, errores y satisfaccion.",
    )


def next_steps(doc):
    add_heading(doc, "8. Recomendaciones para siguiente iteracion", 1)
    for item in [
        "Conectar autenticacion real con backend seguro, HTTPS y validacion del lado servidor.",
        "Sustituir localStorage por API REST o base de datos para usuarios, reservas y lives.",
        "Agregar menu movil compacto para reducir densidad del nav.",
        "Crear archivo editable en Figma/Adobe XD usando la especificacion de frames de este documento.",
        "Ejecutar prueba con usuarios reales y actualizar el reporte con datos medidos.",
    ]:
        add_bullet(doc, item)


def main():
    ensure_dirs()
    doc = setup_document()
    cover(doc)
    scope_matrix(doc)
    wireframes(doc)
    high_fidelity(doc)
    navigation(doc)
    style_guide(doc)
    code_and_api(doc)
    usability(doc)
    next_steps(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
