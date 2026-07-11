from pathlib import Path
import math

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
OUT_DIR = ROOT / "outputs"
ICON_DIR = OUT_DIR / "tripline_doc_icons"
IMAGE_DIR = OUT_DIR / "tripline_doc_images"
DOCX_PATH = OUT_DIR / "TripLine_Iconografia_Imagenes.docx"


COLORS = {
    "neon": RGBColor(57, 255, 20),
    "red_live": RGBColor(255, 32, 32),
    "web_black": RGBColor(2, 2, 2),
    "web_text": RGBColor(232, 232, 224),
    "green": RGBColor(11, 61, 46),
    "green_2": RGBColor(47, 107, 79),
    "gold": RGBColor(217, 164, 65),
    "ink": RGBColor(32, 33, 36),
    "muted": RGBColor(86, 96, 110),
    "light": "EEF5EF",
    "line": "D9E4DC",
    "gold_hex": "D9A441",
    "green_hex": "0B3D2E",
}

FONTS = {
    "display": "Bebas Neue",
    "hud": "Share Tech Mono",
    "body": "Rajdhani",
    "cond": "Barlow Condensed",
    "emoji": "Segoe UI Emoji",
}


def ensure_dirs():
    OUT_DIR.mkdir(exist_ok=True)
    ICON_DIR.mkdir(exist_ok=True)
    IMAGE_DIR.mkdir(exist_ok=True)


def set_run_font(run, name=None, size=None, color=None, bold=None, italic=None):
    if name is None:
        name = FONTS["body"]
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E4DC", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_paragraph(paragraph, before=0, after=6, line_spacing=1.1, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    if align is not None:
        paragraph.alignment = align


def add_text_paragraph(doc, text, size=11, color=None, bold=False, italic=False, align=None, before=0, after=6, font=None):
    p = doc.add_paragraph()
    style_paragraph(p, before=before, after=after, align=align)
    run = p.add_run(text)
    set_run_font(run, font or FONTS["body"], size=size, color=color or COLORS["ink"], bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        style_paragraph(p, before=14, after=8, line_spacing=1.08)
        run = p.add_run(text)
        set_run_font(run, FONTS["display"], 22, COLORS["green"], bold=True)
    else:
        style_paragraph(p, before=8, after=5, line_spacing=1.08)
        run = p.add_run(text)
        set_run_font(run, FONTS["hud"], 11.5, COLORS["green_2"], bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, before=3, after=10, line_spacing=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    set_run_font(r, FONTS["hud"], 8.5, COLORS["muted"], italic=True)


def add_rule(paragraph, color="D9E4DC", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def make_icon(name, path):
    img = Image.new("RGBA", (220, 220), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    green = (11, 61, 46, 255)
    green_light = (238, 245, 239, 255)
    gold = (217, 164, 65, 255)
    muted = (86, 96, 110, 255)

    draw.rounded_rectangle((18, 18, 202, 202), radius=34, fill=green_light, outline=green, width=5)
    draw.ellipse((154, 24, 186, 56), fill=gold)

    if name == "typography":
        draw.line((58, 65, 162, 65), fill=green, width=10)
        draw.line((110, 65, 110, 160), fill=green, width=10)
        draw.line((80, 160, 140, 160), fill=green, width=10)
    elif name == "images":
        draw.rounded_rectangle((54, 62, 166, 154), radius=8, outline=green, width=8)
        draw.polygon([(65, 142), (96, 105), (117, 128), (137, 100), (158, 142)], fill=gold)
        draw.ellipse((132, 76, 150, 94), fill=muted)
    elif name == "icons":
        for x, y in [(78, 82), (142, 82), (78, 146), (142, 146)]:
            draw.ellipse((x - 20, y - 20, x + 20, y + 20), outline=green, width=7)
        draw.line((98, 82, 122, 82), fill=gold, width=6)
        draw.line((78, 102, 78, 126), fill=gold, width=6)
        draw.line((142, 102, 142, 126), fill=gold, width=6)
        draw.line((98, 146, 122, 146), fill=gold, width=6)
    elif name == "code":
        draw.line((88, 82, 58, 110), fill=green, width=9)
        draw.line((58, 110, 88, 138), fill=green, width=9)
        draw.line((132, 82, 162, 110), fill=green, width=9)
        draw.line((162, 110, 132, 138), fill=green, width=9)
        draw.line((118, 72, 100, 148), fill=gold, width=8)
    elif name == "location":
        draw.ellipse((78, 54, 142, 118), outline=green, width=8)
        draw.ellipse((98, 74, 122, 98), fill=gold)
        draw.polygon([(82, 106), (138, 106), (110, 164)], fill=green)
    elif name == "student":
        draw.polygon([(54, 92), (110, 62), (166, 92), (110, 122)], fill=green)
        draw.rectangle((78, 112, 142, 144), fill=gold)
        draw.line((166, 92, 166, 136), fill=green, width=6)
        draw.ellipse((158, 136, 174, 152), fill=green)
    else:
        draw.ellipse((66, 66, 154, 154), outline=green, width=8)
        draw.line((110, 72, 110, 148), fill=gold, width=8)
        draw.line((72, 110, 148, 110), fill=gold, width=8)

    img.save(path)


def create_icons():
    icons = {}
    for name in ("typography", "images", "icons", "code", "location", "student"):
        path = ICON_DIR / f"{name}.png"
        make_icon(name, path)
        icons[name] = path
    return icons


def image_display_size(path, max_width=6.1, max_height=3.55):
    with Image.open(path) as img:
        w, h = img.size
    ratio = w / h
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return width, height


def normalized_image_path(path):
    clean_name = "".join(ch if ch.isalnum() else "_" for ch in path.stem).strip("_")
    target = IMAGE_DIR / f"{clean_name}.png"
    if not target.exists() or target.stat().st_mtime < path.stat().st_mtime:
        with Image.open(path) as img:
            if img.mode not in ("RGB", "RGBA"):
                img = img.convert("RGBA" if "A" in img.getbands() else "RGB")
            img.save(target, "PNG")
    return target


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONTS["body"]
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONTS["body"])
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONTS["body"])
    normal.font.size = Pt(11)
    normal.font.color.rgb = COLORS["ink"]
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("Proyecto TripLine | Junio 2026")
    set_run_font(footer_run, FONTS["hud"], 8.5, COLORS["muted"])

    return doc


def build_cover(doc):
    add_text_paragraph(
        doc,
        "Ingeniería en Desarrollo y Gestión de Software",
        size=15,
        color=COLORS["green_2"],
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=16,
        after=20,
        font=FONTS["cond"],
    )

    p_icon = doc.add_paragraph()
    p_icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_icon = p_icon.add_run("📍")
    set_run_font(r_icon, FONTS["emoji"], 36, COLORS["green"], bold=False)
    style_paragraph(p_icon, before=0, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)

    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=8, line_spacing=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run('Act # - "Proyecto TripLine"')
    set_run_font(r, FONTS["display"], 34, COLORS["green"], bold=True)

    add_text_paragraph(
        doc,
        "Grupo: B7",
        size=14,
        color=COLORS["muted"],
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=4,
        after=32,
        font=FONTS["hud"],
    )

    meta = [
        "Marcos Emmanuel Martínez Alvarado",
        "Matricula: 24007",
        "SANTA CATARINA, N, L, JUNIO 2026",
    ]
    for i, text in enumerate(meta):
        add_text_paragraph(
            doc,
            text,
            size=13 if i < 2 else 11.5,
            color=COLORS["ink"] if i < 2 else COLORS["muted"],
            bold=(i == 0),
            align=WD_ALIGN_PARAGRAPH.CENTER,
            before=0,
            after=8 if i < 2 else 22,
            font=FONTS["body"] if i < 2 else FONTS["hud"],
        )

    hero = ASSETS / "huasteca.jpg"
    if hero.exists():
        hero_for_docx = normalized_image_path(hero)
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(hero_for_docx), width=Inches(5.85))
        style_paragraph(p_img, before=18, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_caption(doc, "Imagen de portada: assets/huasteca.jpg")

    rule = doc.add_paragraph()
    style_paragraph(rule, before=8, after=0)
    add_rule(rule, COLORS["gold_hex"], "12")
    doc.add_page_break()


def build_typography_section(doc):
    add_heading(doc, "Tipografías del código", 1)
    add_text_paragraph(
        doc,
        "El proyecto carga sus fuentes desde Google Fonts y las organiza mediante variables CSS declaradas en :root.",
        size=11,
        color=COLORS["ink"],
        after=8,
    )

    rows = [
        ("--font-display", "Bebas Neue", "Títulos principales, hero, secciones y nombres visuales grandes."),
        ("--font-body", "Rajdhani", "Texto base del sitio y contenido general."),
        ("--font-hud", "Share Tech Mono", "HUD, estados, métricas, badges, controles y textos técnicos."),
        ("--font-cond", "Barlow Condensed", "Subtítulos, textos condensados y apoyo visual."),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [2.1, 1.8, 2.6])
    headers = ["Variable CSS", "Tipografía", "Uso en el código"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, COLORS["light"])
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, after=0)
        r = p.add_run(text)
        set_run_font(r, FONTS["hud"], 9.5, COLORS["green"], bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, (cell, text) in enumerate(zip(cells, row)):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line_spacing=1.1)
            r = p.add_run(text)
            font = FONTS["hud"] if idx == 0 else (text if idx == 1 else FONTS["body"])
            set_run_font(r, font, 10.5, COLORS["ink"], bold=(idx == 0))
    doc.add_paragraph()


def build_iconography_section(doc):
    add_heading(doc, "Iconografía del código", 1)
    add_text_paragraph(
        doc,
        "El proyecto no usa una librería externa de iconos. La iconografía se construye con caracteres Unicode, emoji, botones de control y formas visuales definidas en CSS.",
        size=11,
        after=8,
    )

    rows = [
        ("●", "Señal activa / estado LIVE", "index.html: indicador de señal activa"),
        ("🔇", "Audio desactivado", "index.html y js/app.js: overlay y botón mute"),
        ("🔊", "Audio activado", "index.html y js/app.js: control de volumen"),
        ("⏸ / ▶", "Pausa y reproducción", "index.html y js/app.js: controles de video"),
        ("⛶", "Pantalla completa", "index.html: botón fullscreen"),
        ("📍", "Ubicación GPS", "index.html: coordenadas y ubicaciones de stream"),
        ("🌿", "Selva / Lacandona", "index.html: stream miniatura"),
        ("🏜️", "Desierto / Zona del Silencio", "index.html: stream miniatura"),
        ("📡", "Signal", "index.html: panel HUD"),
        ("🛰", "GPS activo", "index.html: panel HUD"),
        ("⛰", "Altitud", "index.html: panel HUD"),
        ("🌡", "Temperatura", "index.html: panel HUD"),
        ("❤️", "Heart rate", "index.html: panel HUD"),
        ("📺", "Live feed", "index.html: panel HUD"),
        ("⏱", "Duración", "index.html: tarjetas de destinos"),
        ("→ / ←", "Navegación", "index.html y destinos.html: CTA y volver"),
        ("★", "Paquete Extreme", "index.html: badge de paquete"),
        ("✓", "Confirmación de compra", "checkout.html: status-check"),
        ("▸", "Indicadores pseudo-elemento", "css/styles.css: content"),
        ("👁", "Contador de viewers", "css/styles.css: viewers-badge::before"),
        ("live-dot", "Punto animado de transmisión", "css/styles.css: .live-dot"),
        ("hud-corner", "Esquinas HUD decorativas", "css/styles.css: .hud-corner"),
        ("cursor-ring / cursor-dot", "Cursor personalizado", "css/styles.css y index.html"),
    ]

    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [0.95, 2.35, 3.2])
    headers = ["Icono", "Uso", "Origen en el código"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, COLORS["light"])
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(text)
        set_run_font(r, FONTS["hud"], 9.5, COLORS["green"], bold=True)

    for icon, usage, origin in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_border(cell, "E6EEE8", "6")
            set_cell_margins(cell, top=100, bottom=100, start=100, end=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p_icon = cells[0].paragraphs[0]
        style_paragraph(p_icon, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_icon = p_icon.add_run(icon)
        icon_font = FONTS["emoji"] if any(ord(ch) > 127 for ch in icon) else FONTS["hud"]
        set_run_font(r_icon, icon_font, 16 if len(icon) <= 2 else 8.5, COLORS["green"], bold=True)

        p_usage = cells[1].paragraphs[0]
        style_paragraph(p_usage, after=0, line_spacing=1.05)
        r_usage = p_usage.add_run(usage)
        set_run_font(r_usage, FONTS["body"], 10.3, COLORS["ink"], bold=True)

        p_origin = cells[2].paragraphs[0]
        style_paragraph(p_origin, after=0, line_spacing=1.05)
        r_origin = p_origin.add_run(origin)
        set_run_font(r_origin, FONTS["hud"], 8.3, COLORS["muted"])
    doc.add_page_break()


def build_images_section(doc):
    add_heading(doc, "Imágenes integradas del código", 1)
    add_text_paragraph(
        doc,
        "Se insertaron todas las imágenes encontradas en la carpeta assets del proyecto.",
        size=11,
        after=8,
    )

    image_paths = sorted(
        [p for p in ASSETS.iterdir() if p.suffix.lower() in {".jpg", ".jpeg", ".png"}],
        key=lambda p: p.name.lower(),
    )

    for index, path in enumerate(image_paths, start=1):
        add_heading(doc, f"Imagen {index}", 2)
        image_for_docx = normalized_image_path(path)
        width, height = image_display_size(image_for_docx)
        p_img = doc.add_paragraph()
        style_paragraph(p_img, before=2, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_img.add_run().add_picture(str(image_for_docx), width=Inches(width), height=Inches(height))
        add_caption(doc, f"{path.name} | assets/{path.name}")
        if index not in (len(image_paths),) and index % 2 == 0:
            doc.add_page_break()


def main():
    ensure_dirs()
    doc = setup_document()
    build_cover(doc)
    build_typography_section(doc)
    build_iconography_section(doc)
    build_images_section(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
